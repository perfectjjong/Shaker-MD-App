#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""GPC 검증기 공통 — 원본 Accrual 엑셀을 **SSOT 분류(계정·세부채널)** 로 다시 읽는다.
2026-09-11 빌더가 (1) 세부채널 4분할 OR/IR_Main/IR_Others/SME (2) 비B2C 계정 제외(meta.excluded)
(3) 계정 없는 조정행 안분 (4) IR_Others 계정명 노출 로 바뀐 뒤 verify_all/render/matrix 가 전부 낡아
가짜 불일치(335·47·630건)를 내던 것을 여기 한 곳으로 모은다 (2026-09-12).
빌더의 **집계 코드는 재사용하지 않는다**(같은 버그를 공유하지 않기 위함) — 분류 함수만 정의 그 자체이므로 가져온다."""
import collections, glob, re, sys
import openpyxl

sys.path.insert(0, "/home/ubuntu/Shaker-MD-App/docs/dashboards/gpc")
sys.path.insert(0, "/home/ubuntu/2026/10. Automation/03. Operation/00. GPC/_engine")
from build_gpc_dashboard import account_of, row_month, row_cat           # 분류(SSOT 래퍼)만 — row_cat = is_part 부속→Accessory/Others 포함
from gpc_core import sub_channel, account_name, SUB_CHANNELS

ACC_DIR = "/home/ubuntu/2026/02. Operation Team/01. GPC Management/01. Monthly"
# Notes 시트 기준 컬럼(0-idx): M=Cost(12) P=DSI(15) S=Qty(18) T=Value(19) V=YED(21) W=ADC(22) X=VSP(23) Z=EVPD(25) AB=Inv(27)
AMT = {"qty": 18, "gsv": 19, "yed": 21, "adc": 22, "vpd": 25, "dsi": 15, "cogs": 12, "inv": 27, "vsp": 23}
KEYS = list(AMT)


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def latest_accrual():
    return sorted(glob.glob(f"{ACC_DIR}/GPC_Accrual*.xlsx"))[-1]


def load_rows(src=None):
    """→ (rows, excluded, lump_info)
    rows: 대시보드 모집단. dict(y,m,ch,sub,ac,cat, 9지표) — 계정 없는 조정행은 빌더 규칙대로
          같은 (y,m,ch,cat) 의 세부채널×계정 셀에 |GSV| 비중으로 안분해 **이미 더해져 있다**.
    excluded: {(y,m): {지표: 합}} — SSOT 밖(비B2C) 계정. 원본 Σ = rows Σ + excluded Σ 이어야 한다.
    lump_info: (건수, GSV, VSP)"""
    src = src or latest_accrual()
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    cells = collections.defaultdict(lambda: [0.0] * len(KEYS))   # (y,m,ch,sub,ac,cat)
    lump = collections.defaultdict(lambda: [0.0] * len(KEYS))    # (y,m,ch,cat)
    excluded = collections.defaultdict(lambda: {k: 0.0 for k in KEYS})
    for sh in wb.sheetnames:
        mo = re.fullmatch(r"Raw (20\d\d)", sh)
        if not mo:
            continue
        y = int(mo.group(1))
        for row in wb[sh].iter_rows(min_row=2, values_only=True):
            if row[28] is None and row[19] is None and row[12] is None:
                continue
            m = row_month(row)
            chan = str(row[29] or "")
            ch = "IR" if "IR" in chan else ("OR" if "OR" in chan else None)
            if not m or not ch:
                continue
            ac = account_of(row)
            try:
                cid = int(float(row[4]))
            except (TypeError, ValueError):
                cid = None
            sub = sub_channel(cid, ac)
            vals = [num(row[AMT[k]]) for k in KEYS]
            cat = row_cat(row)
            if sub is None and cid is None:                     # 계정 없는 조정행 → 안분 대상
                lk = (y, m, ch, cat)
                for i, v in enumerate(vals):
                    lump[lk][i] += v
                continue
            if sub is None:                                     # SSOT 밖 계정 → 제외
                for k, v in zip(KEYS, vals):
                    excluded[(y, m)][k] += v
                continue
            ch = "OR" if sub == "OR" else "IR"                  # 채널은 SSOT 소속을 따른다
            if ac == "Others" and sub == "IR_Others":
                ac = account_name(cid, row[5])
            a = cells[(y, m, ch, sub, ac, cat)]
            for i, v in enumerate(vals):
                a[i] += v
    wb.close()
    n_lump = sum(1 for v in lump.values() if any(abs(x) > 0 for x in v))
    for (y, m, ch, cat), lv in lump.items():
        cands = [k for k in cells if k[0] == y and k[1] == m and k[2] == ch and k[5] == cat]
        if not cands:
            cands = [k for k in cells if k[0] == y and k[1] == m and k[2] == ch]
        if not cands:
            k = (y, m, ch, "OR" if ch == "OR" else "IR_Others", "Others", cat)
            for i in range(len(KEYS)):
                cells[k][i] += lv[i]
            continue
        tot = sum(abs(cells[k][1]) for k in cands)
        for k in cands:
            w = abs(cells[k][1]) / tot if tot else 1.0 / len(cands)
            for i in range(len(KEYS)):
                cells[k][i] += lv[i] * w
    rows = [dict(y=k[0], m=k[1], ch=k[2], sub=k[3], ac=k[4], cat=k[5], **dict(zip(KEYS, v)))
            for k, v in cells.items()]
    lump_info = (n_lump, sum(v[1] for v in lump.values()), sum(v[8] for v in lump.values()))
    return rows, dict(excluded), lump_info


if __name__ == "__main__":
    rows, ex, li = load_rows()
    print(f"rows {len(rows):,} · excluded 월 {len(ex)} · 조정행 {li[0]}셀 GSV {li[1]:,.0f} VSP {li[2]:,.0f}")
